"""Core logic: read a spreadsheet of multi-series X/Y data and build a PowerPoint
slide with a native, editable XY scatter chart (one colored line per series).

Supported layouts
-----------------
Long  (recommended) - one row per measurement:
    Temperature (°C) | Voltage (V) | Current (mA)
Wide  - one shared X column, then one Y column per series:
    Voltage (V) | 25 °C | 50 °C | 75 °C
"""

from __future__ import annotations

import io

import pandas as pd
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

# Color-blind-friendly palette (Okabe-Ito), distinct in print and on projectors.
PALETTE = ["0072B2", "D55E00", "009E73", "CC79A7", "E69F00", "56B4E9", "000000", "F0E442"]
MARKERS = [
    XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.SQUARE, XL_MARKER_STYLE.TRIANGLE,
    XL_MARKER_STYLE.DIAMOND, XL_MARKER_STYLE.X, XL_MARKER_STYLE.PLUS,
    XL_MARKER_STYLE.STAR, XL_MARKER_STYLE.DASH,
]


# ---------------------------------------------------------------- reading ---
def read_table(file, sheet_name=0) -> pd.DataFrame:
    """Read an .xlsx/.xls/.csv file (path or file-like) into a DataFrame."""
    name = getattr(file, "name", str(file)).lower()
    if name.endswith(".csv"):
        df = pd.read_csv(file)
    else:
        df = pd.read_excel(file, sheet_name=sheet_name)
    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = [str(c).strip() for c in df.columns]
    return df


def list_sheets(file) -> list[str]:
    name = getattr(file, "name", str(file)).lower()
    if name.endswith(".csv"):
        return ["(csv)"]
    return pd.ExcelFile(file).sheet_names


def guess_layout(df: pd.DataFrame) -> str:
    """Guess 'long' if there's a low-cardinality column that could be a label."""
    if df.shape[1] == 3:
        nunique = df.nunique()
        if nunique.min() <= max(10, len(df) // 3):
            return "long"
    return "wide"


def guess_long_columns(df: pd.DataFrame) -> tuple[str, str, str]:
    """Return (group, x, y) guesses: group = column with fewest unique values."""
    cols = list(df.columns)
    group = min(cols, key=lambda c: df[c].nunique())
    rest = [c for c in cols if c != group]
    return group, rest[0], rest[1] if len(rest) > 1 else rest[0]


def _label(value, suffix: str) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f"{value}{suffix}"


def series_from_long(df, group_col, x_col, y_col, label_suffix="") -> list[dict]:
    out = []
    data = df[[group_col, x_col, y_col]].dropna()
    for key, g in data.groupby(group_col, sort=True):
        g = g.sort_values(x_col)
        out.append({
            "name": _label(key, label_suffix),
            "x": pd.to_numeric(g[x_col], errors="coerce").tolist(),
            "y": pd.to_numeric(g[y_col], errors="coerce").tolist(),
        })
    return _clean(out)


def series_from_wide(df, x_col, y_cols, label_suffix="") -> list[dict]:
    out = []
    for c in y_cols:
        g = df[[x_col, c]].dropna().sort_values(x_col)
        out.append({
            "name": f"{c}{label_suffix}",
            "x": pd.to_numeric(g[x_col], errors="coerce").tolist(),
            "y": pd.to_numeric(g[c], errors="coerce").tolist(),
        })
    return _clean(out)


def _clean(series: list[dict]) -> list[dict]:
    """Drop points that aren't numbers, and drop empty series."""
    cleaned = []
    for s in series:
        pts = [(x, y) for x, y in zip(s["x"], s["y"]) if pd.notna(x) and pd.notna(y)]
        if pts:
            cleaned.append({"name": s["name"], "x": [p[0] for p in pts], "y": [p[1] for p in pts]})
    if not cleaned:
        raise ValueError("No numeric data found in the selected columns.")
    return cleaned


# --------------------------------------------------------------- building ---
def build_pptx(
    series: list[dict],
    title: str,
    x_title: str,
    y_title: str,
    legend_title: str = "",
    smooth: bool = True,
    show_markers: bool = True,
    widescreen: bool = True,
) -> bytes:
    """Return the bytes of a one-slide .pptx with a native XY scatter chart."""
    prs = Presentation()
    if widescreen:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    sw, sh = prs.slide_width, prs.slide_height

    # Slide title
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), sw - Inches(1.0), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    p.font.size, p.font.bold = Pt(28), True
    p.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)

    # Chart data: each series keeps its own X values. Shorter series are padded
    # with blank cells so every block in the embedded sheet has the same length;
    # PowerPoint doesn't need this, but LibreOffice/Keynote otherwise drop series.
    cd = XyChartData()
    n_max = max(len(s["x"]) for s in series)
    for s in series:
        ser = cd.add_series(s["name"])
        for x, y in zip(s["x"], s["y"]):
            ser.add_data_point(x, y)
        for _ in range(n_max - len(s["x"])):
            ser.add_data_point(None, None)

    ctype = XL_CHART_TYPE.XY_SCATTER_SMOOTH if smooth else XL_CHART_TYPE.XY_SCATTER_LINES
    if not show_markers:
        ctype = (XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS if smooth
                 else XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS)

    top = Inches(1.2)
    gf = slide.shapes.add_chart(ctype, Inches(0.5), top, sw - Inches(1.0), sh - top - Inches(0.3), cd)
    chart = gf.chart
    chart.font.size = Pt(14)
    chart.font.name = "Calibri"

    # Legend
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(14)

    # Axes with titles and light gridlines
    for axis, text in ((chart.category_axis, x_title), (chart.value_axis, y_title)):
        axis.has_title = True
        axis.axis_title.text_frame.text = text
        f = axis.axis_title.text_frame.paragraphs[0].font
        f.size, f.bold = Pt(16), True
        axis.has_major_gridlines = True
        axis.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        axis.major_gridlines.format.line.width = Pt(0.75)
        axis.tick_labels.font.size = Pt(12)
        axis.format.line.color.rgb = RGBColor(0x59, 0x59, 0x59)

    # Series colors and markers
    for i, plot_ser in enumerate(chart.plots[0].series):
        color = RGBColor.from_string(PALETTE[i % len(PALETTE)])
        line = plot_ser.format.line
        line.color.rgb = color
        line.width = Pt(2.25)
        plot_ser.smooth = smooth
        if show_markers:
            m = plot_ser.marker
            m.style = MARKERS[i % len(MARKERS)]
            m.size = 7
            m.format.fill.solid()
            m.format.fill.fore_color.rgb = color
            m.format.line.color.rgb = color

    # Fixed layout: plot area on the left ~85%, legend pinned top-right, so the
    # legend title label below always sits directly above the legend.
    _set_manual_layout(chart._chartSpace.chart.plotArea, 0.07, 0.03, 0.76, 0.84, inner=True)
    _set_manual_layout(chart._chartSpace.chart.legend, 0.87, LEGEND_Y, 0.11, min(0.85, 0.075 * len(series) + 0.03))

    # Legend title: PowerPoint charts have none, so add a label just above it.
    if legend_title:
        frame_w, frame_h = gf.width, gf.height
        lt = slide.shapes.add_textbox(
            gf.left + int(frame_w * 0.845), gf.top + int(frame_h * LEGEND_Y) - Inches(0.42),
            int(frame_w * 0.155), Inches(0.4))
        lp = lt.text_frame.paragraphs[0]
        lp.text = legend_title
        lp.alignment = PP_ALIGN.CENTER
        lp.font.size, lp.font.bold = Pt(14), True

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


LEGEND_Y = 0.10  # top of the legend, as a fraction of the chart frame height


def _set_manual_layout(el, x, y, w, h, inner=False):
    """Give a chart element (plotArea or legend) a fixed position, in fractions
    of the chart frame. python-pptx has no API for this, so edit the XML."""
    from pptx.oxml.ns import qn
    from lxml import etree

    old = el.find(qn("c:layout"))
    if old is not None:
        el.remove(old)
    layout = etree.SubElement(el, qn("c:layout"))
    ml = etree.SubElement(layout, qn("c:manualLayout"))
    if inner:
        etree.SubElement(ml, qn("c:layoutTarget")).set("val", "inner")
    for tag, val in (("xMode", "edge"), ("yMode", "edge"), ("x", x), ("y", y), ("w", w), ("h", h)):
        etree.SubElement(ml, qn(f"c:{tag}")).set("val", str(val))
    # Schema order: plotArea -> layout is the first child;
    # legend -> layout comes right after legendPos / legendEntry.
    el.remove(layout)
    if el.tag == qn("c:plotArea"):
        el.insert(0, layout)
    else:
        anchor = [c for c in el if c.tag in (qn("c:legendPos"), qn("c:legendEntry"))]
        el.insert(el.index(anchor[-1]) + 1 if anchor else 0, layout)
